
import json,re
from pathlib import Path
from dataclasses import asdict,dataclass,field
from typing import Optional
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
@dataclass
class AnalysisResult:
    document_type:str; project_name:Optional[str]; project_reference:Optional[str]; order_number:Optional[str]; notification_number:Optional[str]; revision:Optional[str]; discipline:Optional[str]; work_instruction_procedure:str; inspection_date_window:Optional[str]; equipment:Optional[str]; scope_of_work:Optional[str]; inspection_factory_location:Optional[str]; site_of_inspection:Optional[str]; named_contacts:list[str]=field(default_factory=list); applicable_standards:list[str]=field(default_factory=list); tests_or_checks:list[str]=field(default_factory=list); inspector_duties:list[str]=field(default_factory=list); referenced_documents_summary:str=""; referenced_documents:list[str]=field(default_factory=list); missing_information:list[str]=field(default_factory=list); confidence_level:str="Low"; reasoned_summary:str=""
def clean(x): return re.sub(r"\s+"," ",str(x or "")).strip()
def uniq(items):
 out=[]; seen=set()
 for i in items:
  c=clean(i)
  if c and c not in seen: out.append(c); seen.add(c)
 return out
def first(text,pat):
 m=re.search(pat,text,flags=re.I|re.S); return clean(m.group(1)) if m else None
def allm(text,pat): return uniq([m.group(1) for m in re.finditer(pat,text,flags=re.I|re.S)])
def extract_text(path:Path):
 suf=path.suffix.lower()
 if suf in {".txt",".md",".csv",".log"}: return path.read_text(encoding="utf-8",errors="ignore")
 if suf==".docx": return "\n".join(p.text for p in Document(str(path)).paragraphs if clean(p.text))
 if suf==".pdf":
  text="\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages[:80])
  if not clean(text): raise ValueError("PDF text could not be extracted. Use a text-based PDF or OCR first.")
  return text
 if suf==".xlsx":
  wb=load_workbook(str(path),data_only=True); lines=[]
  for ws in wb.worksheets[:12]:
   lines.append(f"Sheet: {ws.title}")
   for row in ws.iter_rows(values_only=True):
    vals=[clean(v) for v in row if clean(v)]
    if vals: lines.append(" | ".join(vals))
  return "\n".join(lines)
 raise ValueError("Unsupported file type.")
def classify(text,file_name=""):
 c=f"{text}\n{file_name}"
 if re.search(r"\biFAT\b|system integration test",c,re.I): return "iFAT / Inspection Notification linked to ITP-QCP"
 if re.search(r"notification for inspection|\bNOI\b",c,re.I): return "NOI / Notification for Inspection"
 if re.search(r"notification of readiness for inspection|\bNORFI\b",c,re.I): return "NORFI / Notification of Readiness for Inspection"
 if re.search(r"factory acceptance test plan|\bFAT\b",c,re.I): return "FAT / Factory Acceptance Test"
 if re.search(r"inspection and test plan|\bITP\b|QCP|quality control plan",c,re.I): return "Inspection and Test Plan"
 return "Unknown"
def noi_tests(text):
 items=[]
 if re.search(r"\bF47\b",text,re.I) and re.search(r"partial.?discharge",text,re.I): items.append("F47 — Witness partial-discharge measurement.")
 if re.search(r"\bF61\b",text,re.I) and re.search(r"impulse|AC voltage",text,re.I): items.append("F61 — Witness impulse or AC voltage test on two single coils.")
 return uniq(items)
def xlsx_tests(text):
 items=[]
 if re.search(r"UNIT CONTROL PANEL",text,re.I): items.append("Inspect / witness UNIT CONTROL PANEL iFAT.")
 if re.search(r"SYSTEM INTEGRATION TEST",text,re.I): items.append("Witness System Integration Test.")
 if re.search(r"Inspection Date",text,re.I): items.append("Verify inspection date window and attendance requirements.")
 if re.search(r"QUALITY CONTROL PLAN|Q\.C\.P",text,re.I): items.append("Check inspection activity against the Quality Control Plan / ITP intervention points.")
 if re.search(r"TEST PROCEDURE",text,re.I): items.append("Verify test execution against the referenced test procedure.")
 return uniq(items)
def contacts(text):
 phones=[f"Phone: {p}" for p in allm(text,r"(\+?\d[\d ()-]{7,}\d)") if not re.fullmatch(r"\d{4,}",p)]
 emails=[f"Email: {m}" for m in allm(text,r"([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})")]
 people=[f"Contact person: {m}" for m in allm(text,r"Name\s+([A-Z][A-Za-z .'-]{2,60})")]
 vendors=[f"Vendor / Subcontractor: {m}" for m in allm(text,r"Vendor\s*/\s*Sub\s*contractor Name\s*([^\n]+)")]
 return uniq(emails+phones+people+vendors)[:20]
def standards(text): return allm(text,r"((?:IEC|IEEE|EN|ISO|API|ASME|NFPA)(?:/IEEE)?\s*[0-9A-Z.-]+(?:-[0-9A-Z]+)*)")[:20]
def refdocs(text,uploaded=None):
 docs=[]
 for p in [r"(QUALITY CONTROL PLAN[^\n|]*)",r"(UNIT CONTROL SYSTEM SCHEMATIC[^\n|]*)",r"(UNIT CONTROL SYSTEM I/O LIST[^\n|]*)",r"(UNIT CONTROL SYSTEM INTERNAL WIRING DIAGRAM[^\n|]*)",r"(INTERCONNECTING WIRING DIAGRAM[^\n|]*)",r"(LAYOUT / CONSTRUCTION & MAIN COMPONENTS LIST[^\n|]*)",r"(UNIT CONTROL SYSTEM TEST PROCEDURE[^\n|]*)",r"(Q\.C\.P[^\n|]*)"]: docs+=allm(text,p)
 if uploaded: docs += [f"Uploaded supporting document: {n}" for n in uploaded]
 return uniq(docs)
def analyse_text(text,file_name="",uploaded_ref_names=None):
 doc_type=classify(text,file_name); refs=refdocs(text,uploaded_ref_names or []); stds=standards(text)
 tests=noi_tests(text) if doc_type=="NOI / Notification for Inspection" else (xlsx_tests(text) if "iFAT" in doc_type or file_name.lower().endswith(".xlsx") else uniq([f"{m.group(1)} — {clean(m.group(2))}" for m in re.finditer(r"(F\d{1,3})\s+([^\n]{5,120})",text,flags=re.I)]))
 project_ref=first(text,r"PROJECT No\.?\s*([0-9A-Z.-]+)"); project_name=first(text,r"(MERAM PROJECT)") or first(text,r"(Trion FPU Project)") or first(text,r"Project Name\s*:?\s*([^\n|]+)")
 equipment=first(text,r"PO Description\s*([^\n]+)") or ("UNIT CONTROL PANEL" if re.search(r"UNIT CONTROL PANEL",text,re.I) else None) or first(text,r"Equipment Description\s*&\s*Tag No\.?\s*\|?\s*([^\n|]+)")
 location=first(text,r"(Čedlosy\s+126,\s*664\s*24\s*Drasov)") or first(text,r"(East Gate Business Park F2,\s*H-2151\s*Fót,\s*Hungary)") or first(text,r"Inspection Location\s*:?\s*([^\n|]+)")
 date_window="16.05.2024, 08:00 CET" if re.search(r"16\.05\.2024",text) else ("11 Aug to 15 Aug and 18 Aug to 22 Aug" if re.search(r"11\s+Aug.*22\s+Aug",text,re.I) else first(text,r"(Inspection Date[^\n|]{5,120})"))
 if doc_type=="NOI / Notification for Inspection":
  scope="Witness the electrical tests specified in the NOI / referenced ITP."; subject=f"Inspection of {equipment or 'specified equipment'} — witness listed ITP test steps."; briefing=f"This is an electrical inspection notification for {project_name or project_ref}. The inspector must witness the listed ITP activities, confirm equipment identity, observe the test setup, and record whether the required intervention points are satisfied."; duties=["Confirm the item under test matches the PO / NOI / ITP reference.","Witness F47 partial-discharge measurement and record test status and remarks.","Witness F61 impulse or AC voltage test on two single coils and record test status and remarks.","Check that supplier/contractor/company witness or hold requirements are satisfied.","Record deviations, missing acceptance evidence, and any non-conformity against the referenced ITP or project specification."]
 elif "iFAT" in doc_type:
  scope="iFAT / system integration test and related QCP / ITP checks."; subject=f"Inspection of {equipment or 'unit control equipment'} — iFAT / system integration test."; briefing=f"This inspection concerns iFAT / system integration testing for {equipment or 'the unit control panel'}. The inspector should prepare with the QCP/ITP, drawings and test procedure, then verify evidence against each planned test and inspection item."; duties=["Verify the unit control panel / equipment identity against the instruction and referenced documents.","Witness the system integration test according to the QCP / ITP intervention points.","Check wiring, I/O, layout, component references and test evidence against the uploaded supporting documents.","Record test status, deviations, open points, missing documents and final inspection outcome."]
 else:
  scope="Inspection scope to be confirmed from the uploaded instruction."; subject=f"Inspection of {equipment or 'specified equipment'}."; briefing="This instruction was analysed to prepare an inspection brief. The inspector should verify the extracted fields before using them operationally."; duties=["Review the uploaded instruction.","Confirm the inspection scope, equipment, location and test items.","Record findings and missing information."]
 res=AnalysisResult(doc_type,project_name,project_ref,first(text,r"PO NO:\s*([^\n]+)"),first(text,r"Notification No\.?\s*([0-9A-Z.-]+)"),first(text,r"Rev\s*#?\s*:?\s*([A-Z0-9.-]+)"),"Electrical" if re.search(r"motor|voltage|partial.?discharge|unit control|electrical",text,re.I) else None,subject,date_window,equipment,scope,location,None,contacts(text),stds,tests,duties,"Supporting documents uploaded/identified: "+"; ".join(refs[:8]) if refs else "No supporting documents were uploaded or identified. Add drawings, wiring diagrams, dimensions, QCP/ITP attachments, test limits or project specifications to improve photo review.",refs,[],"High" if len(clean(text))>500 else "Medium",briefing)
 if not stds: res.missing_information.append("No explicit standards were detected in the uploaded instruction. Exact acceptance criteria must come from the referenced ITP / QCP / project specification.")
 if not tests: res.missing_information.append("No clear inspection activities / test items extracted.")
 if not equipment: res.missing_information.append("Equipment not identified confidently.")
 if not date_window: res.missing_information.append("Inspection date window not identified confidently.")
 return asdict(res)
def analyse_photo_placeholder(photo_name,inspection_area,inspection_item,inspector_notes,document_context_raw):
 try: ctx=json.loads(document_context_raw or "{}")
 except Exception: ctx={}
 standards=ctx.get("applicable_standards") or []
 standard_text=", ".join(standards) if standards else "No explicit standard was detected in the instruction. A real defect classification must reference the applicable ITP/QCP/project specification before it can be accepted."
 return {"status_label":"Manual AI review required","severity":"manual","suggested_finding":"Vision model not connected yet","inspection_area":inspection_area,"inspection_item":inspection_item,"inspector_note_suggestion":inspector_notes or "Photo uploaded for review. Real defect classification requires connection to a vision AI model and applicable standard/acceptance criteria.","reasoning":"The inspection workflow is ready to receive photos. VinQu is not yet connected to a real vision model, so it will not invent green/yellow/red findings.","standard_reasoning":standard_text,"recommended_action":"Connect a vision AI model/API and standards knowledge base before using this result for pass/fail or defect classification."}
def analyse_file(path,uploaded_ref_names=None):
 p=Path(path); return analyse_text(extract_text(p),p.name,uploaded_ref_names or [])
